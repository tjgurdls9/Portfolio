# -*- coding: utf-8 -*-
import os #
import requests
from bs4 import BeautifulSoup
import pandas as pd
from datetime import datetime
import time
import re
from openpyxl.styles import PatternFill, Font, Alignment
from google import genai 

# ==========================================
# 1. API KEY SETUP
# ==========================================
MY_API_KEY = "발급받은 API 키를 입력하세요"  
client = genai.Client(api_key=MY_API_KEY)

# 빠른 속도와 넉넉한 한도의 Flash 모델 유지
MODEL_NAME = "gemini-flash-lite-latest"
print(f">> 시스템: 불사조 로직 탑재 완료! [{MODEL_NAME}]로 사냥을 시작합니다.\n")

def ask_ai_master(title, text):
    if not text or len(text) < 100: 
        return "SKIP: 본문 내용 부족"
    
    clean_text = re.sub(r'\s+', ' ', text)[:3500] 
    prompt = f"""
    [임무] 당신은 광주 광산구 소재 '냉동식품 제조사'의 정부지원금 확보를 극대화하는 전략 컨설턴트입니다.
    
    [기업 상세 스펙 - 필수 인지!]
    - 상시근로자 10인 미만의 '제조업 소상공인'입니다.
    - 벤처기업 인증을 보유한 기업입니다.
    - 따라서 '중소기업 지원사업'과 '소상공인 지원사업' 모두 신청 가능합니다.

    [기본 원칙: 긍정적 검토]
    - 우리 회사는 '제조업'이자 '중소기업', '벤처기업'입니다.
    - 식품 전용 사업이 아니더라도, '제조업 전체'나 '중소기업 공통' 지원사업은 우리도 100% 대상이므로 무조건 MATCH 하세요.
    - "이게 될까?" 싶으면 일단 MATCH 하고 봅니다. 놓치는 것보다 과하게 잡는 것이 낫습니다.

    <채택(MATCH) 타겟: 아래 중 하나라도 해당되면 적극적으로 물어올 것>
    1. 자금/금융: 중소기업/소상공인 정책자금, 대출 이자 지원, R&D(연구개발) 지원금, 바우처 사업
    2. 제조/설비: 스마트공장, 생산공정 개선, 자동화/포장 기기 도입, 에너지 효율화 및 탄소중립 지원
    3. 식품/인증: HACCP, FDA, 할랄 등 국내외 식품 인증 획득 비용, 위생/안전 관리, 식약처 관련 지원
    4. 수출/마케팅: 해외 박람회/전시회 참가, 해외 바이어 발굴, 수출 물류비, 카탈로그 제작, 패키지 디자인
    5. 범용 혜택: 벤처기업 우대 혜택, 신규 인력 채용 고용지원금 등
    6. 인건비/채용: 청년 및 중장년 채용 고용지원금, 신규 인력 인건비 보조, 일자리 창출 지원금

    [기각(SKIP) 유의사항]
    - '오직' 특정 업종(예: 조선업체만, IT기업만)으로 한정된 경우에만 SKIP 하세요.
    - 지역이 '광주'를 포함하지 않는 타 지역 전용(예: 경기 전용)일 때만 SKIP 하세요.

    [출력 양식 - 짧고 굵게!]
    - 인사말, 서론, 부연설명 절대 금지. 묻는 말에 답만 하세요.
    - MATCH: [혜택 종류] - [기대 효과 1줄]
    - SKIP: [이유 1줄]

    제목: {title}
    본문: {clean_text}
    """
    
    # 🚨 핵심: 429 에러 시 튕기지 않고 최대 3번까지 60초 대기 후 재시도
    for attempt in range(3):
        try:
            # 질문 속도 여유 있게 조절
            time.sleep(15.0) 
            response = client.models.generate_content(model=MODEL_NAME, contents=prompt)
            res = response.text.strip()
            
            if "MATCH" in res.upper():
                return "MATCH: " + res.replace("MATCH:", "").replace("MATCH", "").strip()
            else:
                return "SKIP: " + res.replace("SKIP:", "").replace("SKIP", "").strip()
                
        except Exception as e:
            error_msg = str(e)
            print(f"\n      >> [진짜 에러 확인]: {error_msg}")
            # 503이랑 UNAVAILABLE(서버 바쁨)까지 싹 다 감시망에 넣습니다!
            if any(x in error_msg for x in ["429", "503", "RESOURCE_EXHAUSTED", "UNAVAILABLE"]):
                print(f"\n      >> [!] 구글 1분 차단기 발동! 60초간 숨 고르고 다시 공격합니다... (현재 {attempt+1}/3회)")
                time.sleep(60.0) # 1분간 완전히 멈춰서 구글 서버가 안정될 때까지 대기
                continue # 버리지 않고 다시 위로 올라가서 재시도
            else:
                return f"ERROR: API 호출 실패 ({error_msg})"
                
    return "ERROR: 과부하로 인해 3회 재시도했으나 실패 (건너뜀)"

def run_immortal_hunter():
    headers = {'User-Agent': 'Mozilla/5.0'}
    final_results = []
    forbidden_pattern = re.compile(r"\[(서울|부산|대구|인천|대전|울산|세종|경기|강원|충북|충남|경북|경남|제주)\]|\((서울|부산|대구|인천|대전|울산|세종|경기|강원|충북|충남|경북|경남|제주)\)")
    
    for page in range(1, 21):
        print(f"--- [Page {page}] Deep Scanning ---")
        url = f"https://www.bizinfo.go.kr/web/lay1/bbs/S1T122C128/AS/74/list.do?cpage={page}&rows=50"
        
        try:
            res = requests.get(url, headers=headers, timeout=10)
            res.encoding = 'utf-8'
            soup = BeautifulSoup(res.text, 'html.parser')
            rows = soup.find('tbody').find_all('tr')
            
            for row in rows:
                a_tag = row.find('a')
                if not a_tag: continue
                title = a_tag.get_text(strip=True)

                if forbidden_pattern.search(title) and not any(r in title for r in ["광주", "전남", "전북", "전라"]):
                    continue

                link = "https://www.bizinfo.go.kr" + a_tag.get('href', '')
                print(f"   검토 중: {title[:35]}...", end=" ")
                
                try:
                    d_res = requests.get(link, headers=headers, timeout=10)
                    d_soup = BeautifulSoup(d_res.text, 'html.parser')
                    detail_area = d_soup.find('div', class_='board_view') 
                    if not detail_area: detail_area = d_soup 
                    detail = " ".join([p.get_text(strip=True) for p in detail_area.find_all(['p', 'div', 'li'])])
                    
                    analysis = ask_ai_master(title, detail)
                    
                    if analysis.startswith("MATCH:"):
                        print(f"\n      >> [O] 채택: {analysis.replace('MATCH:', '')[:60]}...")
                        final_results.append({'공고명': title, 'AI_전략분석': analysis.replace("MATCH:", "").strip(), '상세링크': link})
                    elif analysis.startswith("SKIP:"):
                        print(f"\n      >> [X] 기각: {analysis.replace('SKIP:', '')[:60]}...")
                    else:
                        print(f"\n      >> [!] {analysis[:60]}")
                        
                except Exception as e: 
                    print(f"\n      >> [!] 크롤링 에러: {str(e)}")
                    continue
        except Exception as e: 
            print(f"\n[!] 페이지 로드 에러: {str(e)}")
            break

    if final_results:
        # 1. 폴더 생성 (없으면 새로 만듬)
        folder_name = "공고분석결과"
        if not os.path.exists(folder_name):
            os.makedirs(folder_name)
            
        # 2. 파일명 설정 (날짜 + 시간 포함해서 폴더 안에 저장)
        current_time = datetime.now().strftime('%Y년 %m월 %d일 공고 확인 목록(%H시%M분)')
        file_name = os.path.join(folder_name, f"{current_time}.xlsx")
        
        df = pd.DataFrame(final_results)
        
        with pd.ExcelWriter(file_name, engine='openpyxl') as writer:
            df.to_excel(writer, index=False)
            ws = writer.sheets['Sheet1']
            
            # --- 엑셀 서식 튜닝 (색상 및 너비) ---
            header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
            header_font = Font(color='FFFFFF', bold=True, size=11)
            
            # 제목줄 스타일
            for cell in ws["1:1"]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 열 너비 조절
            ws.column_dimensions['A'].width = 60 # 공고명
            ws.column_dimensions['B'].width = 80 # AI 분석내용
            ws.column_dimensions['C'].width = 90 # 링크
            
            # 내용 정렬 (연한 배경색 추가)
            row_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    cell.fill = row_fill
            # ---------------------------------------

        print(f"\n[작업 완료] 총 {len(final_results)}건 포획!")
        print(f">> 저장 폴더: {os.path.abspath(folder_name)}")
        print(f">> 파일 이름: {os.path.basename(file_name)}")
    else:
        print("\n[작업 완료] 조건에 맞는 공고가 0건입니다.")
if __name__ == "__main__":
    run_immortal_hunter()